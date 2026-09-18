"""Build the 10-slide HieraSync AI deck for github.com/Atulgupta07/Hiera_Sync."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
import os

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "HieraSync_AI_10_Slides.pptx")

# ---------------- palette ----------------
NAVY   = RGBColor(0x0B, 0x1E, 0x3B)
NAVY2  = RGBColor(0x14, 0x33, 0x63)
INDIGO = RGBColor(0x4F, 0x46, 0xE5)
CYAN   = RGBColor(0x06, 0xB6, 0xD4)
GOLD   = RGBColor(0xF5, 0x9E, 0x0B)
LIGHT  = RGBColor(0xF3, 0xF6, 0xFC)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
DARK   = RGBColor(0x1E, 0x29, 0x3B)
MUTED  = RGBColor(0x64, 0x74, 0x8B)
GREEN  = RGBColor(0x10, 0xB9, 0x81)
RED    = RGBColor(0xEF, 0x44, 0x44)
SOFT   = RGBColor(0xEE, 0xF2, 0xFF)
BORDER = RGBColor(0xCB, 0xD5, 0xE1)
MINT   = RGBColor(0xEC, 0xFD, 0xF5)
CREAM  = RGBColor(0xFF, 0xF7, 0xE6)

TOTAL = 10
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


# ---------------- helpers ----------------
def set_bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def new_slide(bg=LIGHT):
    s = prs.slides.add_slide(BLANK)
    set_bg(s, bg)
    return s


def add_shape(slide, shape_type, l, t, w, h, fill=None, line=None,
              line_w=1.0, radius=None):
    shp = slide.shapes.add_shape(shape_type, Inches(l), Inches(t),
                                 Inches(w), Inches(h))
    shp.shadow.inherit = False
    if fill is not None:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    else:
        shp.fill.background()
    if line is not None:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    if radius is not None:
        try:
            shp.adjustments[0].value = radius
        except Exception:
            pass
    return shp


def tf_of(shp, anchor=MSO_ANCHOR.TOP, wrap=True, m=0.08):
    tf = shp.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(m)
    tf.margin_right = Inches(m)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    return tf


def para(tf, text="", size=12, bold=False, color=DARK, align=PP_ALIGN.LEFT,
         first=False, space_after=4, space_before=0, italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    p.alignment = align
    if text:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = "Calibri"
        r.font.italic = italic
    return p


def mixed(tf, parts, align=PP_ALIGN.LEFT, first=False, space_after=4, size=12):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space_after)
    p.alignment = align
    for part in parts:
        text, bold, color = part[0], part[1], part[2]
        sz = part[3] if len(part) > 3 else size
        r = p.add_run()
        r.text = text
        r.font.size = Pt(sz)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = "Calibri"
    return p


def bullets(tf, items, size=12, color=DARK, bullet="▪  ", gap=5,
            bcolor=INDIGO):
    started = False
    for it in items:
        p = tf.paragraphs[0] if not started else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.alignment = PP_ALIGN.LEFT
        if isinstance(it, tuple):
            head, rest = it
            seq = [(bullet, True, bcolor), (head, True, NAVY), (rest, False, color)]
        else:
            seq = [(bullet, True, bcolor), (it, False, color)]
        for txt, b, c in seq:
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(size)
            r.font.bold = b
            r.font.color.rgb = c
            r.font.name = "Calibri"
        started = True


def connector(slide, x1, y1, x2, y2, color=INDIGO, w=2.0, arrow=True,
              ctype=MSO_CONNECTOR.STRAIGHT):
    conn = slide.shapes.add_connector(ctype, Inches(x1), Inches(y1),
                                      Inches(x2), Inches(y2))
    conn.shadow.inherit = False
    conn.line.color.rgb = color
    conn.line.width = Pt(w)
    if arrow:
        try:
            ln = conn.line._get_or_add_ln()
            for child in ln.findall(qn("a:tailEnd")):
                ln.remove(child)
            ln.append(ln.makeelement(qn("a:tailEnd"),
                                     {"type": "triangle", "w": "med", "len": "med"}))
        except Exception:
            pass
    return conn


def chip(slide, l, t, w, h, text, size=9.5, color=NAVY, bg=WHITE,
         line=BORDER, bold=False):
    shp = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h,
                    fill=bg, line=line, radius=0.4)
    tf = tf_of(shp, anchor=MSO_ANCHOR.MIDDLE, m=0.04)
    para(tf, text, size=size, bold=bold, color=color, align=PP_ALIGN.CENTER,
         first=True, space_after=0)
    return shp


def card(slide, l, t, w, h, title, accent=INDIGO, title_size=13,
         fill=WHITE, tcolor=None):
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h,
              fill=fill, line=BORDER, radius=0.06)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, l, t, 0.085, h,
              fill=accent, radius=0.5)
    tb = add_shape(slide, MSO_SHAPE.RECTANGLE, l + 0.22, t + 0.1, w - 0.42, 0.36,
                   fill=None)
    para(tf_of(tb, m=0.0), title, size=title_size, bold=True,
         color=tcolor or NAVY, first=True, space_after=0)
    inner = add_shape(slide, MSO_SHAPE.RECTANGLE, l + 0.22, t + 0.5,
                      w - 0.44, h - 0.62, fill=None)
    return tf_of(inner, m=0.0)


def stat(slide, l, t, w, h, value, lbl, color=INDIGO, bg=WHITE, vsize=21):
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h, fill=bg,
              line=BORDER, radius=0.1)
    tb = add_shape(slide, MSO_SHAPE.RECTANGLE, l + 0.06, t + 0.08, w - 0.12,
                   h - 0.16, fill=None)
    tf = tf_of(tb, anchor=MSO_ANCHOR.MIDDLE, m=0.02)
    para(tf, value, size=vsize, bold=True, color=color, align=PP_ALIGN.CENTER,
         first=True, space_after=1)
    para(tf, lbl, size=9, color=MUTED, align=PP_ALIGN.CENTER, space_after=0)


def header(slide, kicker, title, num):
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 1.18, fill=NAVY)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 1.18, 13.333, 0.045, fill=GOLD)
    tb = add_shape(slide, MSO_SHAPE.RECTANGLE, 0.45, 0.1, 11.0, 1.0, fill=None)
    tf = tf_of(tb, m=0.0)
    para(tf, kicker.upper(), size=10.5, bold=True, color=CYAN, first=True,
         space_after=1)
    para(tf, title, size=25, bold=True, color=WHITE, space_after=0)
    pill = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 12.0, 0.3, 0.9, 0.56,
                     fill=NAVY2, line=CYAN, line_w=1.25, radius=0.45)
    tfp = tf_of(pill, anchor=MSO_ANCHOR.MIDDLE, m=0.02)
    para(tfp, f"{num} / {TOTAL}", size=11.5, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, first=True, space_after=0)


def footer(slide, left="HieraSync AI  •  Smart College Workflow & Event Management System",
           right="github.com/Atulgupta07/Hiera_Sync"):
    tb = add_shape(slide, MSO_SHAPE.RECTANGLE, 0.45, 7.05, 12.43, 0.32, fill=None)
    tf = tf_of(tb, m=0.0)
    p = tf.paragraphs[0]
    p.space_after = Pt(0)
    for txt in (left + "        ", right):
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(8.5)
        r.font.color.rgb = MUTED
        r.font.name = "Calibri"


def content_slide(kicker, title, num):
    s = new_slide(LIGHT)
    header(s, kicker, title, num)
    footer(s)
    return s


def notes(slide, text):
    try:
        slide.notes_slide.placeholders[1].text = text
    except Exception:
        pass


def style_table(table, col_widths, header_fill=NAVY, font_size=10,
                header_size=10.5, row_h=None):
    table.horz_banding = False
    table.first_row = False
    for i, w in enumerate(col_widths):
        table.columns[i].width = Inches(w)
    if row_h:
        for r in table.rows:
            r.height = Inches(row_h)
    for r in range(len(table.rows)):
        for c in range(len(table.columns)):
            cell = table.cell(r, c)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.07)
            cell.margin_right = Inches(0.07)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = header_fill
            else:
                cell.fill.fore_color.rgb = WHITE if r % 2 == 1 else SOFT
            for p in cell.text_frame.paragraphs:
                p.space_after = Pt(1)
                p.space_before = Pt(1)
                for run in p.runs:
                    run.font.size = Pt(header_size if r == 0 else font_size)
                    run.font.bold = (r == 0) or (c == 0 and r > 0)
                    run.font.color.rgb = WHITE if r == 0 else DARK
                    run.font.name = "Calibri"
            try:
                tcPr = cell._tc.get_or_add_tcPr()
                for edge in ("L", "R", "T", "B"):
                    tag = qn(f"a:ln{edge}")
                    for child in tcPr.findall(tag):
                        tcPr.remove(child)
                    ln = tcPr.makeelement(tag, {"w": "12700"})
                    solid = tcPr.makeelement(qn("a:solidFill"), {})
                    solid.append(tcPr.makeelement(qn("a:srgbClr"), {"val": "CBD5E1"}))
                    ln.append(solid)
                    tcPr.append(ln)
            except Exception:
                pass


def make_table(slide, l, t, w, h, data, col_widths, **kw):
    gf = slide.shapes.add_table(len(data), len(data[0]), Inches(l), Inches(t),
                                Inches(w), Inches(h))
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            gf.table.cell(r, c).text = val
    style_table(gf.table, col_widths, **kw)
    return gf


# ================================================================
# SLIDE 1 — TITLE
# ================================================================
s = new_slide(NAVY)
add_shape(s, MSO_SHAPE.OVAL, 10.4, -2.3, 5.2, 5.2, fill=NAVY2)
add_shape(s, MSO_SHAPE.OVAL, -2.1, 4.4, 4.8, 4.8, fill=NAVY2)
add_shape(s, MSO_SHAPE.RECTANGLE, 0.6, 0.5, 1.5, 0.06, fill=GOLD)

pill = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 0.78, 6.9, 0.48,
                 fill=None, line=CYAN, line_w=1.5, radius=0.5)
para(tf_of(pill, anchor=MSO_ANCHOR.MIDDLE, m=0.15),
     "MAJOR PROJECT  •  SMART CAMPUS AUTOMATION  •  SEMINAR PRESENTATION",
     size=10.5, bold=True, color=CYAN, first=True, space_after=0)

tb = add_shape(s, MSO_SHAPE.RECTANGLE, 0.55, 1.42, 9.2, 2.5, fill=None)
tf = tf_of(tb, m=0.0)
para(tf, "HieraSync AI", size=54, bold=True, color=WHITE, first=True,
     space_after=2)
para(tf, "Smart College Workflow, Task & Event Management System", size=20,
     bold=True, color=CYAN, space_after=10)
para(tf, "One platform to assign work, approve requests, track progress and\n"
         "auto-generate reports — with an AI assistant built in.",
     size=13.5, color=RGBColor(0xC7, 0xD2, 0xFE), space_after=0)

link = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 4.02, 5.3, 0.46,
                 fill=NAVY2, line=CYAN, line_w=1.0, radius=0.3)
para(tf_of(link, anchor=MSO_ANCHOR.MIDDLE, m=0.12),
     "Repository:  github.com/Atulgupta07/Hiera_Sync", size=11, bold=True,
     color=WHITE, first=True, space_after=0)

# stats strip
stats = [("20", "API Modules"), ("95", "REST Endpoints"), ("~18,650", "Lines of Code"),
         ("10", "User Roles"), ("17", "App Screens"), ("9", "Data Models")]
x = 0.6
for i, (v, l) in enumerate(stats):
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 4.78, 1.92, 0.92,
              fill=NAVY2, line=RGBColor(0x2A, 0x4A, 0x80), radius=0.12)
    tbx = add_shape(s, MSO_SHAPE.RECTANGLE, x + 0.05, 4.84, 1.82, 0.8, fill=None)
    tfx = tf_of(tbx, anchor=MSO_ANCHOR.MIDDLE, m=0.02)
    para(tfx, v, size=19, bold=True, color=GOLD, align=PP_ALIGN.CENTER,
         first=True, space_after=0)
    para(tfx, l, size=8.5, color=RGBColor(0xA5, 0xB4, 0xFC),
         align=PP_ALIGN.CENTER, space_after=0)
    x += 2.03

# presented by
tb = add_shape(s, MSO_SHAPE.RECTANGLE, 0.6, 5.95, 6.6, 1.1, fill=None)
tf = tf_of(tb, m=0.0)
mixed(tf, [("Presented by:  ", True, CYAN, 11.5),
           ("Atul Gupta  &  Team", False, WHITE, 11.5)], first=True, space_after=3)
mixed(tf, [("Guide:  ", True, CYAN, 11.5),
           ("Prof. ______________", False, WHITE, 11.5)], space_after=3)
mixed(tf, [("Department:  ", True, CYAN, 11.5),
           ("Computer Science & Engineering (AI & ML)", False, WHITE, 11.5)],
      space_after=0)

tb = add_shape(s, MSO_SHAPE.RECTANGLE, 8.0, 6.35, 4.85, 0.6, fill=None)
para(tf_of(tb, m=0.0), "React 19  •  FastAPI  •  Firebase  •  Google Gemini",
     size=11, bold=True, color=GOLD, align=PP_ALIGN.RIGHT, first=True,
     space_after=0)

notes(s, "Good morning. Our project is HieraSync AI — a smart college workflow, "
         "task and event management system. In one line: it replaces notices, "
         "registers, Excel sheets and WhatsApp chasing with a single role-based "
         "web platform that also has an AI assistant. It is a full working system — "
         "20 backend modules, 95 REST endpoints, about 18,650 lines of code, "
         "17 screens and 10 user roles, built on React 19, FastAPI, Firebase and Gemini.")

# ================================================================
# SLIDE 2 — INTRODUCTION
# ================================================================
s = content_slide("Slide 2", "Introduction", 2)

box = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.45, 1.42, 12.43, 0.78,
                fill=SOFT, line=INDIGO, radius=0.08)
tf = tf_of(box, anchor=MSO_ANCHOR.MIDDLE, m=0.18)
mixed(tf, [("HieraSync AI ", True, NAVY, 13),
           ("(backend service name: ", False, DARK, 12.5),
           ("CampusPulse API", True, INDIGO, 12.5),
           (") is a web-based platform that brings every department activity — tasks, approvals, "
            "events, goals, files and reports — into ", False, DARK, 12.5),
           ("one secure, role-based system", True, NAVY, 12.5),
           (", with an AI assistant that answers questions from live data.", False, DARK, 12.5)],
      first=True, space_after=0)

tf = card(s, 0.45, 2.34, 7.55, 2.52, "What the system actually does", INDIGO)
bullets(tf, [
    ("Single login for the full hierarchy — ", "Principal → HOD → Faculty → Staff / Students (10 roles)"),
    ("Work is assigned as tasks — ", "with deadline, priority, progress %, comments & attachments"),
    ("Requests & approvals move in a chain — ", "HOD stage, then Principal stage, fully logged"),
    ("AI assistant (Google Gemini) — ", "answers using your real tasks, approvals and events"),
    ("Automatic reminders — ", "in-app + WhatsApp, checked by a background job every 5 minutes"),
    ("Academic calendar PDFs — ", "uploaded once, parsed by AI into categorised events"),
], size=11.5, gap=6)

cards = [("95", "REST\nEndpoints", INDIGO), ("20", "Backend\nModules", CYAN),
         ("17", "Frontend\nScreens", GREEN), ("9", "Firestore\nData Models", GOLD)]
x = 8.2
for i, (v, l, c) in enumerate(cards):
    cx = x + (i % 2) * 2.4
    cy = 2.34 + (i // 2) * 1.3
    stat(s, cx, cy, 2.25, 1.16, v, l, color=c, vsize=20)

tf = card(s, 0.45, 4.98, 12.43, 1.95, "From manual work  →  to one automated system", GOLD)
y = 5.62
steps_old = ["Notice on board", "WhatsApp group", "Register / Excel", "Phone follow-up", "Report typed by hand"]
steps_new = ["Login (role-based)", "Assign task", "Auto notify + remind", "Track progress live", "One-click report"]
label_w, chip_w, gap = 1.28, 1.92, 0.24

lab = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.68, y, label_w, 0.48,
                fill=RGBColor(0xFE, 0xE2, 0xE2), line=RED, radius=0.3)
para(tf_of(lab, anchor=MSO_ANCHOR.MIDDLE, m=0.03), "BEFORE", size=9.5, bold=True,
     color=RED, align=PP_ALIGN.CENTER, first=True, space_after=0)
x = 0.68 + label_w + gap
for i, t in enumerate(steps_old):
    chip(s, x, y, chip_w, 0.48, t, size=9, color=DARK, bg=WHITE, line=BORDER)
    if i < len(steps_old) - 1:
        connector(s, x + chip_w, y + 0.24, x + chip_w + gap, y + 0.24, color=MUTED, w=1.5)
    x += chip_w + gap

y2 = 6.24
lab = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.68, y2, label_w, 0.48,
                fill=MINT, line=GREEN, radius=0.3)
para(tf_of(lab, anchor=MSO_ANCHOR.MIDDLE, m=0.03), "WITH HIERASYNC", size=8,
     bold=True, color=GREEN, align=PP_ALIGN.CENTER, first=True, space_after=0)
x = 0.68 + label_w + gap
for i, t in enumerate(steps_new):
    chip(s, x, y2, chip_w, 0.48, t, size=9, color=NAVY, bg=MINT, line=GREEN, bold=True)
    if i < len(steps_new) - 1:
        connector(s, x + chip_w, y2 + 0.24, x + chip_w + gap, y2 + 0.24, color=GREEN, w=1.5)
    x += chip_w + gap

notes(s, "HieraSync AI is a single web platform for department operations. The backend "
         "service is internally called CampusPulse API. Everything a department does — "
         "assigning tasks, approving requests, scheduling events, setting goals, storing "
         "files and making reports — happens inside one role-based system. "
         "The bottom strip shows the core idea: today the flow is notice → WhatsApp → "
         "register → phone follow-up → hand-typed report. With HieraSync it becomes "
         "login → assign → auto-notify → track → one-click report. Same work, but recorded, "
         "automatic and measurable.")

# ================================================================
# SLIDE 3 — PROBLEM STATEMENT & OBJECTIVES
# ================================================================
s = content_slide("Slide 3", "Problem Statement & Objectives", 3)

box = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.45, 1.42, 12.43, 0.72,
                fill=CREAM, line=GOLD, radius=0.08)
tf = tf_of(box, anchor=MSO_ANCHOR.MIDDLE, m=0.18)
mixed(tf, [("Problem Statement:  ", True, GOLD, 12.5),
           ("College departments still run on notices, WhatsApp groups, registers and Excel sheets — "
            "so work is invisible, approvals are slow, deadlines are missed and every report has to be "
            "rebuilt by hand.", False, DARK, 12.5)], first=True, space_after=0)

tf = card(s, 0.45, 2.28, 6.1, 4.62, "Problems identified", RED)
bullets(tf, [
    "No single place to see who is doing what, and by when",
    "Approvals travel on paper or WhatsApp — no status, no trail",
    "Deadlines are missed because every reminder is manual",
    "The academic calendar is a PDF that nobody re-reads",
    "NAAC / NBA reports are re-created by hand every time",
    "Faculty workload is invisible, so work is unevenly divided",
    "Files, proofs and discussions are scattered across chats",
    "New faculty have no record of past activities",
], size=11.5, gap=9, bcolor=RED)

tf = card(s, 6.78, 2.28, 6.1, 4.62, "Objectives of the project", GREEN)
bullets(tf, [
    ("1. ", "Build one secure, role-based platform for the whole hierarchy"),
    ("2. ", "Digitise task assignment with deadline, priority, progress & proof"),
    ("3. ", "Automate a two-stage approval workflow (HOD → Principal)"),
    ("4. ", "Send reminders automatically — in-app and on WhatsApp"),
    ("5. ", "Convert uploaded academic-calendar files into events using AI"),
    ("6. ", "Provide an AI assistant that answers from live department data"),
    ("7. ", "Warn about risky deadlines before they are actually missed"),
    ("8. ", "Generate dashboards, analytics and exportable reports in one click"),
], size=11.5, gap=9, bullet="", bcolor=GREEN)

notes(s, "The problem statement in one line: department work today is invisible, slow and "
         "unrecorded. On the left are the eight problems we found. On the right are our eight "
         "objectives, and each objective directly answers a problem — one platform for the "
         "hierarchy, digital tasks, automated approvals, automatic reminders, AI calendar "
         "parsing, an AI assistant, deadline-risk warnings and one-click reports. "
         "Every one of these eight objectives is implemented in the code we will show next.")

# ================================================================
# SLIDE 4 — LITERATURE SURVEY
# ================================================================
s = content_slide("Slide 4", "Literature Survey", 4)

box = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.45, 1.4, 12.43, 0.6,
                fill=SOFT, line=INDIGO, radius=0.1)
para(tf_of(box, anchor=MSO_ANCHOR.MIDDLE, m=0.18),
     "We studied the systems a college department already uses. Each one solves a part of the problem — "
     "none of them covers a department end-to-end.",
     size=12, color=DARK, first=True, space_after=0)

data = [
    ["Existing approach", "What it does well", "Limitation for a department", "HieraSync AI answer"],
    ["Registers, notices &\nWhatsApp groups",
     "Zero cost, everyone already\nknows how to use it",
     "No tracking, no history, no reports;\ninformation is lost in chat",
     "Same simplicity, but every action is\nrecorded, searchable and reportable"],
    ["College ERP / MIS\nsoftware",
     "Strong for admission, fees,\nattendance and exam records",
     "Stores records only — no day-to-day\ntask assignment or approval flow",
     "Adds the missing workflow, approval\nand progress-tracking layer"],
    ["Corporate PM tools\n(Trello, Asana, Jira)",
     "Excellent task boards and\nprogress views",
     "Paid per user; no HOD→Principal\nhierarchy, no academic calendar",
     "Free stack, hierarchy-aware roles and\nacademic-specific modules"],
    ["Google Classroom /\nWorkspace",
     "Good for teaching material\nand file sharing",
     "Built for teacher↔student, not for\nstaff duties, approvals or events",
     "Covers the staff-side operations that\nClassroom does not touch"],
    ["Generic AI chatbots\n(ChatGPT etc.)",
     "Very good natural-language\nanswers",
     "Knows nothing about your tasks,\ndeadlines or pending approvals",
     "Gemini + live Firestore context, so it\nanswers about YOUR actual work"],
]
make_table(s, 0.45, 2.14, 12.43, 4.28, data, [2.35, 2.75, 3.75, 3.58],
           font_size=9.5, header_size=10.5)

box = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.45, 6.5, 12.43, 0.45,
                fill=MINT, line=GREEN, radius=0.2)
mixed(tf_of(box, anchor=MSO_ANCHOR.MIDDLE, m=0.18),
      [("Research gap:  ", True, GREEN, 11.5),
       ("no single free system combines hierarchy-based task workflow + two-stage approvals + "
        "academic calendar intelligence + an AI assistant that reads live institutional data.",
        False, DARK, 11.5)], first=True, space_after=0)

notes(s, "For the literature survey we compared the five things a department could use today. "
         "Manual registers are free but leave no record. College ERPs handle admission, fees and "
         "exams but not day-to-day work. Corporate tools like Trello and Asana have great boards "
         "but charge per user and have no HOD-to-Principal hierarchy. Google Classroom is built for "
         "teacher-student, not staff operations. And generic AI chatbots give good language answers "
         "but know nothing about our actual tasks. The gap at the bottom is exactly what our project "
         "fills: hierarchy workflow + approvals + academic calendar intelligence + a context-aware AI.")

# ================================================================
# SLIDE 5 — SYSTEM DESIGN
# ================================================================
s = content_slide("Slide 5", "System Design — Three-Tier Architecture", 5)

tiers = [
    (0.45, "TIER 1  —  PRESENTATION", "Client / Browser", INDIGO,
     ["React 19 Single Page App", "17 screens • Redux Toolkit store",
      "ProtectedRoute + AuthContext", "FullCalendar • Recharts • dnd-kit",
      "Tailwind CSS 4 responsive UI"]),
    (4.82, "TIER 2  —  APPLICATION", "FastAPI Server", CYAN,
     ["20 routers  →  95 REST endpoints", "JWT verify + check_role() guard",
      "APScheduler background jobs", "Document parser (PDF/DOCX/XLSX)",
      "Risk engine + AI service layer"]),
    (9.19, "TIER 3  —  DATA & SERVICES", "Firebase + External APIs", GREEN,
     ["Firestore NoSQL collections", "Firebase Auth • Firebase Storage",
      "Google Gemini 1.5 Flash API", "WhatsApp Cloud API v22.0",
      "SMTP e-mail notifications"]),
]
for l, kick, name, color, items in tiers:
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, l, 1.5, 3.7, 3.55, fill=WHITE,
              line=color, line_w=1.5, radius=0.06)
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, l, 1.5, 3.7, 0.72, fill=color, radius=0.1)
    tb = add_shape(s, MSO_SHAPE.RECTANGLE, l + 0.12, 1.54, 3.46, 0.64, fill=None)
    tfh = tf_of(tb, anchor=MSO_ANCHOR.MIDDLE, m=0.02)
    para(tfh, kick, size=8.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
         first=True, space_after=1)
    para(tfh, name, size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
         space_after=0)
    y = 2.36
    for it in items:
        chip(s, l + 0.18, y, 3.34, 0.46, it, size=9.5, color=DARK, bg=LIGHT, line=BORDER)
        y += 0.52

connector(s, 4.17, 3.05, 4.79, 3.05, color=NAVY, w=2.5)
connector(s, 4.79, 3.55, 4.17, 3.55, color=MUTED, w=2.0)
connector(s, 8.54, 3.05, 9.16, 3.05, color=NAVY, w=2.5)
connector(s, 9.16, 3.55, 8.54, 3.55, color=MUTED, w=2.0)
for lx, t1, t2 in [(4.16, "HTTPS", "REST / JSON"), (8.53, "Admin", "SDK / HTTPS")]:
    tb = add_shape(s, MSO_SHAPE.RECTANGLE, lx, 2.52, 0.68, 0.45, fill=None)
    tfl = tf_of(tb, anchor=MSO_ANCHOR.MIDDLE, m=0.0)
    para(tfl, t1, size=7.5, bold=True, color=NAVY, align=PP_ALIGN.CENTER,
         first=True, space_after=0)
    para(tfl, t2, size=6.5, color=MUTED, align=PP_ALIGN.CENTER, space_after=0)

# request flow
add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.45, 5.2, 12.43, 0.85, fill=WHITE,
          line=BORDER, radius=0.1)
tb = add_shape(s, MSO_SHAPE.RECTANGLE, 0.6, 5.28, 1.5, 0.7, fill=None)
tfl = tf_of(tb, anchor=MSO_ANCHOR.MIDDLE, m=0.0)
para(tfl, "REQUEST", size=9, bold=True, color=NAVY, first=True, space_after=0)
para(tfl, "FLOW", size=9, bold=True, color=NAVY, space_after=0)
flow = ["1  User clicks in\nReact screen", "2  fetch() with\nBearer JWT",
        "3  FastAPI checks\ntoken + role", "4  Firestore read\n/ write",
        "5  Notify • WhatsApp\n• Gemini", "6  JSON back →\nUI updates"]
x = 1.95
for i, t in enumerate(flow):
    shp = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 5.32, 1.63, 0.62,
                    fill=SOFT, line=INDIGO, radius=0.15)
    para(tf_of(shp, anchor=MSO_ANCHOR.MIDDLE, m=0.03), t, size=8, color=NAVY,
         align=PP_ALIGN.CENTER, first=True, space_after=0)
    if i < len(flow) - 1:
        connector(s, x + 1.63, 5.63, x + 1.76, 5.63, color=INDIGO, w=1.5)
    x += 1.76

# cross-cutting
add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.45, 6.2, 12.43, 0.72, fill=NAVY, radius=0.1)
tb = add_shape(s, MSO_SHAPE.RECTANGLE, 0.6, 6.26, 2.0, 0.6, fill=None)
para(tf_of(tb, anchor=MSO_ANCHOR.MIDDLE, m=0.0), "CROSS-CUTTING", size=9,
     bold=True, color=CYAN, first=True, space_after=0)
cc = ["JWT HS256 · 30 min", "check_role() RBAC", "CORS middleware",
      "APScheduler · every 5 min", "Central logging", "503 handlers for API errors"]
x = 2.35
for t in cc:
    chip(s, x, 6.33, 1.72, 0.46, t, size=8, color=WHITE, bg=NAVY2,
         line=RGBColor(0x2A, 0x4A, 0x80))
    x += 1.78

notes(s, "This is the system design — a classic three-tier architecture. Tier 1 is the browser: "
         "a React 19 single-page app with 17 screens. Tier 2 is the FastAPI server: 20 routers "
         "exposing 95 REST endpoints, plus the JWT check, the role guard, the background scheduler, "
         "the document parser and the AI service. Tier 3 is Firebase — Firestore for data, Auth for "
         "identity, Storage for files — plus the external Gemini and WhatsApp APIs. "
         "Follow the request flow at the bottom: the user clicks, the browser sends a fetch call with "
         "a Bearer JWT, FastAPI validates the token and the role, reads or writes Firestore, triggers "
         "notifications or AI if needed, and returns JSON that updates the UI. The dark strip lists "
         "what applies to every request — JWT, role checks, CORS, the 5-minute scheduler and logging.")

# ================================================================
# SLIDE 6 — TECHNOLOGY USED
# ================================================================
s = content_slide("Slide 6", "Technology Used", 6)

data = [
    ["Layer", "Technologies used in the project", "Why it was chosen"],
    ["Frontend",
     "React 19.2, TypeScript 6, Vite 8, Tailwind CSS 4, Redux Toolkit, React Router 7,\n"
     "FullCalendar 6, Recharts 3, dnd-kit (drag & drop), Framer Motion, Lucide icons",
     "Fast dev server, type safety,\nready-made calendar & charts"],
    ["Backend",
     "Python + FastAPI, Uvicorn ASGI server, Pydantic v2 schemas, pydantic-settings,\n"
     "modular router architecture (20 routers)",
     "Async, very fast, and gives free\nSwagger API docs at /docs"],
    ["Database",
     "Firebase Firestore (NoSQL document DB) — collections for users, tasks, events,\n"
     "approvals, requests, goals, notifications, join_requests",
     "Real-time, schema-flexible,\nno server to maintain, free tier"],
    ["Security",
     "Firebase Authentication, JWT (python-jose, HS256, 30-min expiry),\n"
     "passlib + bcrypt password hashing, check_role() RBAC dependency",
     "Industry-standard auth; one guard\nprotects every endpoint"],
    ["AI &\nIntegrations",
     "Google Gemini 1.5 Flash REST API with a local fallback engine,\n"
     "WhatsApp Cloud API v22.0, SMTP e-mail, APScheduler background jobs",
     "Real AI answers; system still works\neven if the AI key is missing"],
    ["Files &\nReports",
     "pypdf, python-docx, openpyxl for parsing PDF / Word / Excel calendars;\n"
     "Firebase Storage for attachments; CSV & report export endpoints",
     "Accepts whatever format the office\nalready has"],
    ["Tools &\nDeployment",
     "Git & GitHub, VS Code, Oxlint, Swagger UI; backend → Cloud Run / Render,\n"
     "frontend → Vercel / Firebase Hosting",
     "Free hosting tiers suitable for a\ncollege deployment"],
]
make_table(s, 0.45, 1.4, 12.43, 5.5, data, [1.35, 7.5, 3.58],
           font_size=9.5, header_size=11)

notes(s, "The technology stack, layer by layer. Frontend is React 19 with TypeScript and Vite, "
         "styled with Tailwind 4; FullCalendar gives us the calendar view, Recharts the graphs, "
         "dnd-kit the drag-and-drop task board. Backend is Python FastAPI with Uvicorn and Pydantic — "
         "chosen because it is async, fast, and auto-generates Swagger documentation at slash docs, "
         "which we use for testing. The database is Firebase Firestore, a NoSQL document database, "
         "so there is no SQL server to maintain and it stays on the free tier. Security is Firebase "
         "Auth plus our own JWT with bcrypt hashing and a single role-guard dependency. "
         "For AI we call Google Gemini 1.5 Flash over REST — and importantly we wrote a local "
         "fallback engine, so if the API key is missing the assistant still answers from Firestore data. "
         "WhatsApp Cloud API sends reminders, and pypdf, python-docx and openpyxl let us read academic "
         "calendars in whatever format the office already has.")

# ================================================================
# SLIDE 7 — DEVELOPED MODULES
# ================================================================
s = content_slide("Slide 7", "Developed Modules — 20 Modules, 95 Endpoints", 7)

data = [
    ["Module", "API", "What it does in the system"],
    ["Authentication & Employees", "8",
     "Register, login (JWT), /me profile, forgot-password, employee list & CRUD by role"],
    ["Departments & Join Requests", "10",
     "Create department, unique join code, faculty join request → HOD approve / reject"],
    ["Tasks & Reviews", "6",
     "Create, update, delete, assign & co-assign tasks; submit for review; AI risk score"],
    ["Comments & Attachments", "7",
     "Threaded comments on a task; upload, list, download and delete proof files"],
    ["Task Requests", "4",
     "8 request types (leave, resource, deadline extension…); approval auto-creates a task"],
    ["Two-Stage Approvals", "5",
     "Raise request → HOD stage → Principal stage; approve / reject with comments & trail"],
    ["Events & Smart Calendar", "6",
     "Event CRUD, participants, meeting link, conflict checking, recurring events"],
    ["Institutional Calendar (AI)", "8",
     "Upload PDF/DOCX/XLSX → AI extracts & categorises activities → draft → publish → history"],
    ["Department Goals", "8",
     "Long-term goals with category, target date, status and milestone tracking"],
    ["AI Assistant", "4",
     "Dashboard summary, Gemini chat on live data, AI report generation, checklist suggestions"],
    ["Notifications & WhatsApp", "13",
     "8 notification types, unread count, mark-read; WhatsApp send, opt-in, webhook, dedup"],
    ["Reports & Analytics", "8",
     "Dashboard stats, recent activity, summaries, per-task reports, export, faculty performance"],
    ["Files, Search, Settings", "8",
     "Firebase Storage upload/download, global search, profile & department settings, health check"],
]
make_table(s, 0.45, 1.38, 12.43, 5.06, data, [2.75, 0.62, 9.06],
           font_size=9.5, header_size=10.5)

# bottom enum strip
add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.45, 6.52, 12.43, 0.46, fill=NAVY, radius=0.2)
tb = add_shape(s, MSO_SHAPE.RECTANGLE, 0.6, 6.54, 12.2, 0.42, fill=None)
tfk = tf_of(tb, anchor=MSO_ANCHOR.MIDDLE, m=0.0)
mixed(tfk, [("Roles: ", True, CYAN, 9), ("Admin · Principal · HOD · Teacher · Faculty · TA · Student · Student-Rep · Lab Assistant · Staff    ", False, WHITE, 9),
            ("Task status: ", True, CYAN, 9), ("TODO → IN_PROGRESS → IN_REVIEW → COMPLETED / OVERDUE    ", False, WHITE, 9),
            ("Priority: ", True, CYAN, 9), ("Low · Medium · High · Urgent", False, WHITE, 9)],
      first=True, space_after=0)

notes(s, "These are the modules we have actually built — 20 backend routers giving 95 REST endpoints. "
         "Read the API column: authentication and employees is 8 endpoints, departments and join "
         "requests 10, tasks 6, comments and attachments 7, task requests 4, approvals 5, events 6, "
         "the AI institutional calendar 8, goals 8, the AI assistant 4, notifications and WhatsApp 13, "
         "reports and analytics 8, and files, search and settings 8. "
         "Two modules are our highlights. The Institutional Calendar module lets the office upload the "
         "official academic calendar as a PDF, Word or Excel file; our parser extracts the activities "
         "and the AI categorises them into eleven types like Examination, Workshop or Holiday, we "
         "review it as a draft and then publish it to everyone's calendar. The AI Assistant module "
         "sends the user's live tasks, approvals and events to Gemini as context, so its answers are "
         "about your real work, not generic advice. The strip at the bottom shows the enumerations "
         "defined in our models — ten roles, five task states and four priority levels.")

# ================================================================
# SLIDE 8 — ADVANTAGES & APPLICATIONS
# ================================================================
s = content_slide("Slide 8", "Advantages & Applications", 8)

tf = card(s, 0.45, 1.4, 6.35, 4.5, "Advantages", GREEN)
bullets(tf, [
    "One login replaces registers, Excel sheets and WhatsApp chasing",
    "Every task, approval and comment is time-stamped and searchable",
    "Reminders are automatic — nothing depends on someone remembering",
    "AI risk score warns before a deadline is actually missed",
    "Reports and analytics are produced in one click, not one week",
    "Role-based access: each person sees only what concerns them",
    "Works even without the AI key — local fallback engine answers",
    "Free & open stack — Firebase free tier, no per-user licence cost",
    "Browser-based: works on laptop and mobile, nothing to install",
], size=11, gap=7, bcolor=GREEN)

add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.03, 1.4, 5.85, 4.5, fill=WHITE,
          line=BORDER, radius=0.06)
add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.03, 1.4, 0.085, 4.5, fill=INDIGO, radius=0.5)
tb = add_shape(s, MSO_SHAPE.RECTANGLE, 7.25, 1.5, 5.4, 0.36, fill=None)
para(tf_of(tb, m=0.0), "Applications", size=13, bold=True, color=NAVY,
     first=True, space_after=0)

apps = [("Department task &\nduty allotment", INDIGO), ("Seminar, workshop &\nevent management", CYAN),
        ("Leave, resource &\ndocument approvals", GREEN), ("Academic calendar\npublishing", GOLD),
        ("NAAC / NBA / AICTE\ndocumentation", INDIGO), ("Faculty workload &\nperformance analytics", CYAN),
        ("Student-rep & lab\nstaff coordination", GREEN), ("Any hierarchical office\n(school, admin block)", GOLD)]
for i, (t, c) in enumerate(apps):
    ax = 7.25 + (i % 2) * 2.72
    ay = 1.95 + (i // 2) * 0.96
    shp = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, ax, ay, 2.6, 0.84,
                    fill=LIGHT, line=BORDER, radius=0.12)
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, ax, ay, 0.07, 0.84, fill=c, radius=0.5)
    para(tf_of(shp, anchor=MSO_ANCHOR.MIDDLE, m=0.12), t, size=10,
         bold=True, color=NAVY, align=PP_ALIGN.CENTER, first=True, space_after=0)

add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.45, 6.02, 12.43, 0.92, fill=WHITE,
          line=BORDER, radius=0.1)
tb = add_shape(s, MSO_SHAPE.RECTANGLE, 0.62, 6.1, 1.6, 0.76, fill=None)
para(tf_of(tb, anchor=MSO_ANCHOR.MIDDLE, m=0.0), "WHO\nBENEFITS", size=9.5,
     bold=True, color=NAVY, first=True, space_after=0)
ben = [("Principal", "full-institute visibility & approvals", INDIGO),
       ("HOD", "assign, monitor and control department work", CYAN),
       ("Faculty", "clear list of duties, deadlines & reminders", GREEN),
       ("Staff / Students", "one channel to raise and track requests", GOLD)]
x = 2.3
for name, desc, c in ben:
    shp = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 6.14, 2.56, 0.68,
                    fill=LIGHT, line=c, radius=0.14)
    tfb = tf_of(shp, anchor=MSO_ANCHOR.MIDDLE, m=0.08)
    para(tfb, name, size=10.5, bold=True, color=c, align=PP_ALIGN.CENTER,
         first=True, space_after=1)
    para(tfb, desc, size=8, color=MUTED, align=PP_ALIGN.CENTER, space_after=0)
    x += 2.63

notes(s, "The advantages on the left are the practical gains: one login instead of four tools, "
         "a permanent time-stamped record, automatic reminders, early risk warnings, one-click "
         "reports, role-based privacy, and zero licence cost because the whole stack runs on free "
         "tiers. Note the seventh point — even without a Gemini key the assistant still works, "
         "because we wrote a local fallback engine. "
         "On the right are the application areas — from daily duty allotment and event management "
         "to NAAC and NBA documentation, and the same system fits any hierarchical office, not just "
         "a college department. The bottom strip shows what each role gains: the Principal gets "
         "visibility, the HOD gets control, faculty get clarity, and staff and students get a single "
         "channel to raise and track requests.")

# ================================================================
# SLIDE 9 — REFERENCES
# ================================================================
s = content_slide("Slide 9", "References", 9)

box = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.45, 1.4, 12.43, 0.55,
                fill=SOFT, line=INDIGO, radius=0.1)
para(tf_of(box, anchor=MSO_ANCHOR.MIDDLE, m=0.18),
     "Official documentation of the technologies used to design, build and deploy the system.",
     size=11.5, color=DARK, first=True, space_after=0)

refs_left = [
    ("[1]", "FastAPI Documentation", "fastapi.tiangolo.com", "Backend framework, routers, dependency injection"),
    ("[2]", "React Documentation", "react.dev", "React 19 components, hooks and routing"),
    ("[3]", "Firebase Firestore Docs", "firebase.google.com/docs/firestore", "NoSQL collections, queries and security rules"),
    ("[4]", "Firebase Authentication", "firebase.google.com/docs/auth", "Identity, sign-in methods, Admin SDK"),
    ("[5]", "Google Gemini API", "ai.google.dev/gemini-api/docs", "Gemini 1.5 Flash generateContent REST API"),
]
refs_right = [
    ("[6]", "WhatsApp Cloud API", "developers.facebook.com/docs/whatsapp", "Template messages, webhooks, v22.0"),
    ("[7]", "APScheduler Documentation", "apscheduler.readthedocs.io", "Background jobs for 5-minute reminders"),
    ("[8]", "Pydantic v2 Documentation", "docs.pydantic.dev", "Request/response schemas and validation"),
    ("[9]", "Tailwind CSS & Vite", "tailwindcss.com  •  vite.dev", "Utility-first styling and build tooling"),
    ("[10]", "Project Repository", "github.com/Atulgupta07/Hiera_Sync", "Complete source code of this project"),
]

for col, items in enumerate([refs_left, refs_right]):
    x = 0.45 + col * 6.33
    y = 2.1
    for tag, name, url, why in items:
        shp = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 6.1, 0.88,
                        fill=WHITE, line=BORDER, radius=0.1)
        add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 0.07, 0.88,
                  fill=GOLD if tag == "[10]" else INDIGO, radius=0.5)
        tb = add_shape(s, MSO_SHAPE.RECTANGLE, x + 0.18, y + 0.06, 5.85, 0.76, fill=None)
        tfr = tf_of(tb, anchor=MSO_ANCHOR.MIDDLE, m=0.0)
        mixed(tfr, [(tag + "  ", True, GOLD if tag == "[10]" else INDIGO, 10.5),
                    (name, True, NAVY, 11)], first=True, space_after=1)
        para(tfr, url, size=9.5, color=INDIGO, space_after=1)
        para(tfr, why, size=8.5, color=MUTED, space_after=0)
        y += 0.96

notes(s, "These are our references — the official documentation we used while building each part: "
         "FastAPI and Pydantic for the backend, React and Tailwind for the frontend, Firebase "
         "Firestore, Authentication and Storage for data, the Google Gemini API for the assistant, "
         "the WhatsApp Cloud API for reminders and APScheduler for the background jobs. "
         "Reference ten is our own repository, which holds the complete source code of the system.")

# ================================================================
# SLIDE 10 — CONCLUSION
# ================================================================
s = content_slide("Slide 10", "Conclusion", 10)

box = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.45, 1.4, 12.43, 0.72,
                fill=SOFT, line=INDIGO, radius=0.08)
para(tf_of(box, anchor=MSO_ANCHOR.MIDDLE, m=0.18),
     "HieraSync AI converts scattered, manual department work into one recorded, automated and "
     "measurable workflow — and adds an AI layer that reads the institution's own live data.",
     size=13, bold=True, color=NAVY, align=PP_ALIGN.CENTER, first=True, space_after=0)

concl = [
    ("BUILT", "20 backend modules, 95 REST endpoints,\n17 screens, ~18,650 lines of code", INDIGO),
    ("WORKING", "End-to-end flow runs: register → join dept →\nassign task → approve → report", GREEN),
    ("INTELLIGENT", "Gemini assistant, deadline-risk engine and\nAI academic-calendar parsing", CYAN),
    ("PRACTICAL", "Free stack, browser-based, mobile friendly,\nrole-based and deployment ready", GOLD),
]
x = 0.45
for title, body, c in concl:
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.28, 3.0, 1.5, fill=WHITE,
              line=BORDER, radius=0.1)
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.28, 3.0, 0.42, fill=c, radius=0.2)
    tb = add_shape(s, MSO_SHAPE.RECTANGLE, x + 0.1, 2.3, 2.8, 0.38, fill=None)
    para(tf_of(tb, anchor=MSO_ANCHOR.MIDDLE, m=0.0), title, size=11, bold=True,
         color=WHITE, align=PP_ALIGN.CENTER, first=True, space_after=0)
    tb = add_shape(s, MSO_SHAPE.RECTANGLE, x + 0.12, 2.76, 2.76, 0.96, fill=None)
    para(tf_of(tb, anchor=MSO_ANCHOR.MIDDLE, m=0.0), body, size=9.5, color=DARK,
         align=PP_ALIGN.CENTER, first=True, space_after=0)
    x += 3.17

tf = card(s, 0.45, 3.95, 6.35, 2.28, "Future Scope", GOLD)
bullets(tf, [
    "Android / iOS mobile app with push notifications",
    "Direct export into NAAC & NBA report templates",
    "Timetable and attendance system integration",
    "Multi-department and multi-college rollout",
    "Formal unit testing, load testing and security audit",
    "Offline mode with automatic sync when back online",
], size=10.5, gap=4, bcolor=GOLD)

tf = card(s, 7.03, 3.95, 5.85, 2.28, "What we learned", INDIGO)
bullets(tf, [
    "Designing a real REST API with role-based security",
    "Working with a NoSQL database and cloud services",
    "Integrating a live LLM with safe fallback handling",
    "Building a responsive React + TypeScript front end",
    "Turning a real institutional problem into software",
], size=10.5, gap=6)

add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.45, 6.35, 12.43, 0.6, fill=NAVY, radius=0.2)
tb = add_shape(s, MSO_SHAPE.RECTANGLE, 0.6, 6.37, 12.13, 0.56, fill=None)
tfc = tf_of(tb, anchor=MSO_ANCHOR.MIDDLE, m=0.0)
mixed(tfc, [("Thank You", True, GOLD, 15), ("        Questions are welcome.        ", False, WHITE, 12),
            ("github.com/Atulgupta07/Hiera_Sync", False, CYAN, 11)],
      align=PP_ALIGN.CENTER, first=True, space_after=0)

notes(s, "To conclude. We set out to remove manual, invisible department work, and we have built a "
         "working system: 20 modules, 95 endpoints, 17 screens, about 18,650 lines of code. The full "
         "flow runs end to end — a faculty member registers, joins a department by code, receives a "
         "task, submits it for review, the HOD and Principal approve, and the report is generated. "
         "On top of that we added intelligence: a Gemini-powered assistant, a deadline-risk engine "
         "and AI parsing of the official academic calendar. "
         "Future scope is a mobile app, direct NAAC and NBA export, timetable integration, "
         "multi-department rollout and formal testing. Thank you — I am happy to take questions.")

prs.save(OUT)
print(f"Saved: {OUT}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
